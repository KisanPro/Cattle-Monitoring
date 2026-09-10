import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

src_md = sys.argv[1] if len(sys.argv) > 1 else r"F:\JRF\Output\Consolidated Document\CONSOLIDATED_DOCUMENT.md"
out_docx = sys.argv[2] if len(sys.argv) > 2 else r"F:\JRF\Output\Consolidated Document\KisanPro_Technical_User_Manual.docx"
img_dir = sys.argv[3] if len(sys.argv) > 3 else r"F:\JRF\Output\Consolidated Document\images"
doc_title = sys.argv[4] if len(sys.argv) > 4 else "KisanPro Integrated Technical User Manual"

doc = Document()
for sec in doc.sections:
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.8)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    footer = sec.footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_run = f_p.add_run(f"{doc_title} | Page ")
    f_run.font.name = "Calibri"
    f_run.font.size = Pt(9)
    f_run.font.color.rgb = RGBColor(0x7F, 0x8C, 0x8D)
    fldSimple = OxmlElement("w:fldSimple")
    fldSimple.set(qn("w:instr"), "PAGE")
    f_run._r.append(fldSimple)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def sanitize(text):
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1', text)
    return text.strip()

with open(src_md, "r", encoding="utf-8") as f:
    lines = f.readlines()

in_code = False
code_lines = []
in_table = False
table_rows = []

i = 0
while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    
    if stripped.startswith("```"):
        if not in_code:
            in_code = True
            code_lines = []
        else:
            in_code = False
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(code_text)
            r.font.name = "Courier New"
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        i += 1
        continue
    
    if in_code:
        code_lines.append(line.rstrip())
        i += 1
        continue
        
    if "|" in stripped and stripped.startswith("|"):
        if not in_table:
            in_table = True
            table_rows = [stripped]
        else:
            table_rows.append(stripped)
        i += 1
        continue
    elif in_table:
        in_table = False
        parsed_rows = []
        for tr in table_rows:
            if re.match(r'^\|[\s\-:|]+\|$', tr):
                continue
            parsed_rows.append([c.strip() for c in tr.strip('|').split('|')])
        if parsed_rows:
            num_cols = max(len(r) for r in parsed_rows)
            t = doc.add_table(rows=len(parsed_rows), cols=num_cols)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            t.autofit = True
            for r_idx, row_data in enumerate(parsed_rows):
                for c_idx, val in enumerate(row_data):
                    if c_idx < num_cols:
                        cell = t.cell(r_idx, c_idx)
                        cell.text = sanitize(val)
                        p = cell.paragraphs[0]
                        run = p.runs[0] if p.runs else p.add_run(cell.text)
                        if r_idx == 0:
                            set_cell_background(cell, "1B365D")
                            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                            run.font.name = "Calibri"
                            run.font.size = Pt(9)
                            run.font.bold = True
                            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        else:
                            bg = "F7F9FB" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg)
                            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
                            run.font.name = "Calibri"
                            run.font.size = Pt(8.5)
                            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
            doc.add_paragraph()

    img_match = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
    if img_match:
        caption = img_match.group(1)
        rel_path = img_match.group(2)
        img_full = os.path.normpath(os.path.join(img_dir, os.path.basename(rel_path)))
        if os.path.exists(img_full):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run()
            run.add_picture(img_full, width=Inches(3.2))
            c_p = doc.add_paragraph()
            c_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c_p.paragraph_format.space_before = Pt(2)
            c_p.paragraph_format.space_after = Pt(8)
            c_run = c_p.add_run(caption)
            c_run.font.name = "Calibri"
            c_run.font.size = Pt(8.5)
            c_run.font.italic = True
            c_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1
        continue

    if stripped.startswith("# "):
        doc.add_page_break()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(sanitize(stripped[2:]))
        run.font.name = "Calibri"
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    elif stripped.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(sanitize(stripped[3:]))
        run.font.name = "Calibri"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x2B, 0x54, 0x7E)
    elif stripped.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(sanitize(stripped[4:]))
        run.font.name = "Calibri"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x48, 0x63, 0xA0)
    elif stripped.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(sanitize(stripped[5:]))
        run.font.name = "Calibri"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
    elif stripped.startswith("- ") or stripped.startswith("* "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(sanitize(stripped[2:]))
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    elif re.match(r'^\d+\.\s', stripped):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(sanitize(re.sub(r'^\d+\.\s', '', stripped)))
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    elif stripped == "---":
        pass
    elif stripped:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(sanitize(stripped))
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
    i += 1

doc.save(out_docx)
print(f"[SUCCESS] Compiled: {out_docx}")
