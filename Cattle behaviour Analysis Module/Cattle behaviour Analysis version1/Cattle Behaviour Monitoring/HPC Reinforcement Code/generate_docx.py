import os
import sys
import re
import io
import base64
import requests
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from PIL import Image

# Force UTF-8 stdout if possible
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

def set_cell_background(cell, hex_color):
    """Sets background shading on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_page_number_to_footer(footer):
    """Inserts dynamic PAGE field into the document footer."""
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Page ")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(128, 128, 128)
    
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    p._p.append(fldSimple)

def fetch_mermaid_image(mermaid_code: str) -> io.BytesIO:
    """Encodes mermaid syntax to base64 and retrieves rendered image from mermaid.ink."""
    try:
        clean_code = mermaid_code.strip()
        encoded = base64.b64encode(clean_code.encode("utf-8")).decode("ascii")
        url = f"https://mermaid.ink/img/{encoded}?bgColor=FFFFFF"
        response = requests.get(url, timeout=25, verify=False)
        if response.status_code == 200:
            return io.BytesIO(response.content)
    except Exception as e:
        print(f"[!] Warning: Could not fetch Mermaid image: {e}")
    return None

def clean_inline_markdown(text: str) -> str:
    """Removes inline markdown symbols."""
    text = re.sub(r'\*\*(.*?)\*\*', r'\1', text)
    text = re.sub(r'\*(.*?)\*', r'\1', text)
    text = re.sub(r'`(.*?)`', r'\1', text)
    text = re.sub(r'\$(.*?)\$', r'\1', text)
    text = re.sub(r'\[(.*?)\]\((.*?)\)', r'\1', text)
    return text.strip()

def compile_document(md_path: Path, output_base: str = "Kisan_HPC_Consolidated_Documentation.docx"):
    print(f"[*] Reading markdown documentation from {md_path}...")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        add_page_number_to_footer(section.footer)

    # Document Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(40, 40, 40)

    # Split into lines
    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lang = ""
    code_buffer = []
    
    in_table = False
    table_lines = []

    print("[*] Parsing sections and embedding high-res visual diagrams...")

    while i < len(lines):
        line = lines[i]
        
        # Handle Code Blocks & Mermaid Diagrams
        if line.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip("`").strip().lower()
                code_buffer = []
                i += 1
                continue
            else:
                in_code_block = False
                block_content = "\n".join(code_buffer)
                
                if code_lang == "mermaid":
                    print("    -> Rendering Mermaid diagram...")
                    img_stream = fetch_mermaid_image(block_content)
                    if img_stream:
                        try:
                            # Choose width according to diagram type
                            if "classDiagram" in block_content:
                                width = Inches(4.2)
                            elif "sequenceDiagram" in block_content:
                                width = Inches(4.2)
                            elif "graph TD" in block_content:
                                width = Inches(3.4)
                            else:
                                width = Inches(5.2)
                                
                            img_p = doc.add_paragraph()
                            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            img_p.paragraph_format.space_before = Pt(8)
                            img_p.paragraph_format.space_after = Pt(12)
                            doc.add_picture(img_stream, width=width)
                        except Exception as ex:
                            print(f"[!] Failed to insert picture: {ex}")
                    else:
                        # Fallback to monospace code
                        p = doc.add_paragraph()
                        p.paragraph_format.left_indent = Inches(0.4)
                        run = p.add_run(block_content)
                        run.font.name = "Courier New"
                        run.font.size = Pt(8.5)
                else:
                    # Regular Code / Bash / Text Block
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    cell = table.cell(0, 0)
                    set_cell_background(cell, "F5F5F7")
                    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
                    
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run(block_content)
                    run.font.name = "Courier New"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(30, 30, 30)
                    
                    doc.add_paragraph() # Spacing
                
                i += 1
                continue
                
        if in_code_block:
            code_buffer.append(line)
            i += 1
            continue

        # Handle Tables
        if "|" in line and not line.strip().startswith("```"):
            if not in_table:
                in_table = True
                table_lines = [line]
            else:
                table_lines.append(line)
            i += 1
            continue
        else:
            if in_table:
                in_table = False
                # Parse and render table
                rows_data = []
                for tline in table_lines:
                    if re.match(r'^\s*\|?\s*[-:]+[-| :]*\|?\s*$', tline):
                        continue # Separator row
                    cols = [clean_inline_markdown(c) for c in tline.split('|')[1:-1]]
                    if cols:
                        rows_data.append(cols)
                
                if rows_data:
                    num_rows = len(rows_data)
                    num_cols = len(rows_data[0])
                    doc_table = doc.add_table(rows=num_rows, cols=num_cols)
                    doc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                    doc_table.autofit = True
                    
                    for r_idx, row in enumerate(rows_data):
                        for c_idx, cell_value in enumerate(row):
                            if c_idx < len(doc_table.columns):
                                cell = doc_table.cell(r_idx, c_idx)
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_before = Pt(4)
                                p.paragraph_format.space_after = Pt(4)
                                run = p.add_run(cell_value)
                                
                                if r_idx == 0:
                                    set_cell_background(cell, "1B365D") # Navy Header
                                    set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(255, 255, 255)
                                    run.font.size = Pt(10)
                                else:
                                    bg_color = "F9FAFC" if r_idx % 2 == 1 else "FFFFFF"
                                    set_cell_background(cell, bg_color)
                                    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
                                    run.font.size = Pt(9.5)
                                    run.font.color.rgb = RGBColor(50, 50, 50)
                                    
                    doc.add_paragraph() # Spacing
                table_lines = []

        # Headings & Paragraphs
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("# CHAPTER"):
            doc.add_page_break()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_inline_markdown(stripped.lstrip("#").strip()))
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 54, 93) # Deep Navy
        elif stripped.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(24)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_inline_markdown(stripped[2:]))
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(27, 54, 93)
        elif stripped.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_inline_markdown(stripped[3:]))
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(41, 128, 185) # Slate Blue
        elif stripped.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(clean_inline_markdown(stripped[4:]))
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(70, 70, 70)
        elif stripped.startswith("> [!") or stripped.startswith(">"):
            # Callout block
            callout_text = clean_inline_markdown(stripped.lstrip(">").strip())
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "EBF3FB")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(callout_text)
            run.font.size = Pt(10)
            run.font.italic = True
            run.font.color.rgb = RGBColor(30, 70, 120)
        elif stripped.startswith("* ") or stripped.startswith("- "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean_inline_markdown(stripped[2:]))
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            text_part = re.sub(r'^\d+\.\s', '', stripped)
            run = p.add_run(clean_inline_markdown(text_part))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(clean_inline_markdown(stripped))
            
        i += 1

    # Save with lock fallback
    filename = output_base
    version = 2
    while True:
        try:
            doc.save(filename)
            print(f"[+] SUCCESS: Publication-grade Word Document compiled to: {filename}")
            break
        except PermissionError:
            filename = f"{output_base.replace('.docx', '')}_v{version}.docx"
            version += 1

if __name__ == "__main__":
    doc_path = Path("documents/CONSOLIDATED_DOCUMENT.md")
    if not doc_path.exists():
        doc_path = Path("CONSOLIDATED_DOCUMENT.md")
    compile_document(doc_path)
